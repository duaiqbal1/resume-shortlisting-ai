"""
resume_parser.py
----------------
Handles text extraction from PDF, DOCX and TXT resumes, and extracts
structured candidate information (name guess, email, phone, skills,
years of experience) using regex-based rules.
"""

import os
import re

import pdfplumber
import docx


# A reasonably broad skill vocabulary used for automatic skill detection.
# Feel free to extend this list for your domain (data science, marketing, etc.)
SKILL_KEYWORDS = [
    "python", "java", "c++", "c#", "javascript", "typescript", "react", "angular",
    "vue", "node.js", "node", "django", "flask", "fastapi", "spring", "spring boot",
    "sql", "mysql", "postgresql", "mongodb", "redis", "docker", "kubernetes",
    "aws", "azure", "gcp", "git", "github", "linux", "html", "css", "bootstrap",
    "tailwind", "machine learning", "deep learning", "nlp", "tensorflow",
    "pytorch", "scikit-learn", "pandas", "numpy", "data analysis", "power bi",
    "tableau", "excel", "communication", "leadership", "project management",
    "agile", "scrum", "rest api", "graphql", "microservices", "ci/cd",
    "unit testing", "selenium", "android", "ios", "swift", "kotlin", "flutter",
    "php", "laravel", "wordpress", "photoshop", "figma", "ui/ux", "seo",
]

EMAIL_REGEX = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
PHONE_REGEX = re.compile(r"(\+?\d{1,3}[-.\s]?)?\(?\d{3,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}")
EXPERIENCE_REGEX = re.compile(
    r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years|yrs|year)\s*(?:of)?\s*experience",
    re.IGNORECASE,
)


def extract_text_from_file(filepath: str) -> str:
    """Extract raw text from a PDF, DOCX or TXT file."""
    ext = os.path.splitext(filepath)[1].lower()

    try:
        if ext == ".pdf":
            return _extract_from_pdf(filepath)
        elif ext == ".docx":
            return _extract_from_docx(filepath)
        elif ext == ".txt":
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    except Exception as e:
        print(f"[resume_parser] Error extracting {filepath}: {e}")
        return ""

    return ""


def _extract_from_pdf(filepath: str) -> str:
    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_from_docx(filepath: str) -> str:
    document = docx.Document(filepath)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    # Also pull text from tables (skills tables are common in resumes)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def _guess_name(text: str, fallback: str) -> str:
    """Very simple heuristic: assume the first non-empty line that looks like
    a name (2-4 capitalized words, no digits/@ symbols) is the candidate name."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    for line in lines[:8]:
        words = line.split()
        if 1 < len(words) <= 4 and not any(ch.isdigit() for ch in line) and "@" not in line:
            if all(w[0].isupper() for w in words if w[0].isalpha()):
                return line.title()
    return os.path.splitext(fallback)[0].replace("_", " ").replace("-", " ").title()


def _extract_skills(text: str) -> list:
    text_lower = text.lower()
    found = []
    for skill in SKILL_KEYWORDS:
        pattern = r"\b" + re.escape(skill.lower()) + r"\b"
        if re.search(pattern, text_lower):
            found.append(skill)
    return sorted(set(found))


def _extract_experience_years(text: str) -> float:
    matches = EXPERIENCE_REGEX.findall(text)
    if matches:
        try:
            return max(float(m) for m in matches)
        except ValueError:
            return 0.0
    return 0.0


def extract_candidate_info(text: str, original_filename: str) -> dict:
    """Build a structured dictionary describing one candidate."""
    email_match = EMAIL_REGEX.search(text)
    phone_match = PHONE_REGEX.search(text)

    return {
        "filename": original_filename,
        "name": _guess_name(text, original_filename),
        "email": email_match.group(0) if email_match else "Not found",
        "phone": phone_match.group(0) if phone_match else "Not found",
        "skills": _extract_skills(text),
        "experience_years": _extract_experience_years(text),
        "raw_text": text,
    }
